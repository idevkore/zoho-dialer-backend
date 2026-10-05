#!/usr/bin/env python3
"""
Builds docs/market/pricing-model.xlsx: the HaulOS pricing model, with live
formulas on every sheet.

    python3 build_model.py
    python3 build_model.py --set p_dialer=99 --set retention=3 --out /tmp/what-if.xlsx

Inputs live in assumptions.py; --set changes one for a single build without
editing the file. --actuals and --out point at another usage file or workbook. Measured usage is read from
docs/market/inputs/usage-actuals.json; a metric that is null there falls back
to the matching assumption on the UsageActuals sheet, and says so.

Nothing here imports from the app, and the app does not import this.
"""

import argparse
import json
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

from assumptions import SECTIONS

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
ACTUALS = ROOT / 'docs' / 'market' / 'inputs' / 'usage-actuals.json'
OUT = ROOT / 'docs' / 'market' / 'pricing-model.xlsx'

# Kore brand: purple headers on near-black, lime for the numbers that matter.
INK = '0B0B10'
PURPLE = '8051FF'
LIME = 'B9FF00'
LAVENDER = 'EEE8FF'
BLUE = '0000FF'
MUTED = '6B6B76'

FMT = {
    'usd': '$#,##0.00', 'usd0': '$#,##0', 'usd4': '$0.0000', 'usd6': '$0.000000',
    'pct': '0.0%', 'int': '#,##0', 'num': '#,##0.00', 'text': '@',
}
CASES = (('C', 'low'), ('D', 'base'), ('E', 'high'))


# ── Cell helpers ────────────────────────────────────────────────────────────

def put(ws, ref, value, fmt=None, bold=False, color=None, fill=None, italic=False, wrap=False, size=None):
    cell = ws[ref]
    cell.value = value
    if fmt:
        cell.number_format = FMT.get(fmt, fmt)
    cell.font = Font(name='Calibri', bold=bold, italic=italic, color=color, size=size or 10)
    if fill:
        cell.fill = PatternFill('solid', fgColor=fill)
    cell.alignment = Alignment(wrap_text=wrap, vertical='top')
    return cell


def inp(ws, ref, value, fmt=None):
    """An input: blue, as a reader of a financial model expects."""
    return put(ws, ref, value, fmt, color=BLUE)


def key(ws, ref, value, fmt=None):
    """A headline output: lime on near-black."""
    return put(ws, ref, value, fmt, bold=True, color=LIME, fill=INK)


def title(ws, text, sub, ncols):
    for col in range(1, ncols + 1):
        for row in (1, 2):
            ws.cell(row=row, column=col).fill = PatternFill('solid', fgColor=INK)
    put(ws, 'A1', text, bold=True, color=LIME, fill=INK, size=14)
    put(ws, 'A2', sub, color='C9C9D4', fill=INK, size=9)
    ws.row_dimensions[1].height = 22
    ws.sheet_properties.tabColor = PURPLE
    ws.sheet_view.showGridLines = False


def header(ws, row, labels, start=1):
    for i, label in enumerate(labels):
        ref = f'{get_column_letter(start + i)}{row}'
        put(ws, ref, label, bold=True, color=PURPLE, fill=INK, wrap=True)
    ws.row_dimensions[row].height = 30


def section(ws, row, text, ncols):
    for col in range(1, ncols + 1):
        ws.cell(row=row, column=col).fill = PatternFill('solid', fgColor=LAVENDER)
    put(ws, f'A{row}', text, bold=True, color=INK, fill=LAVENDER)


def widths(ws, spec):
    for col, width in spec.items():
        ws.column_dimensions[col].width = width


# ── Assumptions ─────────────────────────────────────────────────────────────

REF = {}        # key → row on the Assumptions sheet
FIRST_ROW = 5


def a(name):
    """A cell on the Assumptions sheet, from any other sheet."""
    return f'Assumptions!$C${REF[name]}'


def assign_rows():
    row = FIRST_ROW
    for _, entries in SECTIONS:
        row += 1                       # the section's own row
        for entry in entries:
            REF[entry['key']] = row
            row += 1
        row += 1                       # a blank row after each section


def build_assumptions(wb):
    ws = wb.active
    ws.title = 'Assumptions'
    title(ws, 'HaulOS pricing model: assumptions',
          'Blue is an input. Black is a formula. Every input carries a source and a confidence. '
          'ASSUMPTION means nobody has measured it.', 7)
    header(ws, 4, ['Key', 'Input', 'Value', 'Unit', 'Source', 'Confidence', 'Note'])
    row = FIRST_ROW
    for name, entries in SECTIONS:
        section(ws, row, name, 7)
        row += 1
        for e in entries:
            put(ws, f'A{row}', e['key'], color=MUTED, size=8)
            put(ws, f'B{row}', e['label'], wrap=True)
            if e['kind'] == 'input':
                inp(ws, f'C{row}', e['value'], e['fmt'])
            else:
                formula = re.sub(r'\{(\w+)\}', lambda m: f'$C${REF[m.group(1)]}', e['formula'])
                put(ws, f'C{row}', f'={formula}', e['fmt'])
            put(ws, f'D{row}', e['unit'], color=MUTED)
            put(ws, f'E{row}', e['source'], wrap=True, size=9)
            conf = e['conf']
            put(ws, f'F{row}', conf, bold=conf == 'ASSUMPTION', color='B00020' if conf == 'ASSUMPTION' else None)
            put(ws, f'G{row}', e['note'], wrap=True, size=9, color=MUTED)
            wb.defined_names[e['key']] = DefinedName(e['key'], attr_text=f"Assumptions!$C${row}")
            row += 1
        row += 1
    widths(ws, {'A': 16, 'B': 52, 'C': 14, 'D': 22, 'E': 58, 'F': 14, 'G': 70})
    ws.freeze_panes = 'C5'
    return ws


# ── UsageActuals ────────────────────────────────────────────────────────────

GROUPS = [
    ('voice', ['calls_out', 'calls_in', 'connected_calls', 'avg_duration_sec', 'median_duration_sec',
               'billed_min_total', 'billed_min_pstn_out', 'billed_min_pstn_in', 'billed_min_browser_leg',
               'recorded_min', 'voicemail_drops', 'amd_calls', 'active_agents', 'peak_concurrent']),
    ('sms', ['msgs_out', 'msgs_in', 'segments_out', 'segments_in', 'delivery_rate', 'opt_outs',
             'filtered_errors', 'carrier_fee_usd']),
    ('zoho_crm', ['licensed_users', 'edition', 'api_credits_per_day_avg', 'api_credits_per_day_peak',
                  'api_credit_limit', 'call_records_created']),
    ('zoho_marketing', ['sends']),
]

# (metric, label, unit, numerator(s), per-seat?, assumption key)
PROFILE = [
    ('out_min', 'Billed minutes, PSTN outbound', 'per seat / month', [('voice', 'billed_min_pstn_out')], True, 'd_out_min'),
    ('in_min', 'Billed minutes, PSTN inbound', 'per seat / month', [('voice', 'billed_min_pstn_in')], True, 'd_in_min'),
    ('browser_min', 'Billed minutes, browser leg', 'per seat / month', [('voice', 'billed_min_browser_leg')], True, 'd_browser_min'),
    ('rec_min', 'Recorded minutes', 'per seat / month', [('voice', 'recorded_min')], True, 'd_rec_min'),
    ('calls', 'Calls, out plus in', 'per seat / month', [('voice', 'calls_out'), ('voice', 'calls_in')], True, 'd_calls'),
    ('amd_calls', 'Answering-machine-detection calls', 'per seat / month', [('voice', 'amd_calls')], True, 'd_amd'),
    ('seg_out', 'Outbound SMS segments', 'per seat / month', [('sms', 'segments_out')], True, 'd_seg_out'),
    ('seg_in', 'Inbound SMS segments', 'per seat / month', [('sms', 'segments_in')], True, 'd_seg_in'),
    ('ai_convos', 'AI conversations', 'per tenant / month', None, False, 'd_ai_convos'),
]

USED = {}       # metric → {'low': ref, 'base': ref, 'high': ref} on UsageActuals
DATA_ROW = 7


def build_usage_actuals(wb):
    ws = wb.create_sheet('UsageActuals')
    doc = json.loads(ACTUALS.read_text())
    window = doc.get('window', {})

    col_of, col = {}, 5
    for group, fields in GROUPS:
        for field in fields:
            col_of[(group, field)] = get_column_letter(col)
            col += 1
    helper_start = col + 1
    last_col = helper_start + len(PROFILE)

    title(ws, 'Usage actuals',
          f"From inputs/usage-actuals.json, pulled {doc.get('pulled_at', '?')}, window {window.get('from', '?')} "
          f"to {window.get('to', '?')}. A blank cell is a null: not retrieved, not zero.", last_col)

    put(ws, 'A4', 'Include = 1 counts the row in the per-seat figures. Development, sandbox and part-months start '
                  'at 0. Set a ramp-up month to 0 to leave it out.', italic=True, color=MUTED, size=9)
    header(ws, 6, ['Tenant', 'Environment', 'Month', 'Include'])
    for (group, field), letter in col_of.items():
        put(ws, f'{letter}5', group, color=MUTED, size=8)
        header(ws, 6, [field], start=ws[f'{letter}6'].column)

    rows = []
    to_month = str(window.get('to', ''))[:7]
    for tenant in doc.get('tenants', []):
        for month in tenant.get('months', []):
            rows.append((tenant, month))

    for i, (tenant, month) in enumerate(rows):
        r = DATA_ROW + i
        put(ws, f'A{r}', tenant.get('tenant'))
        put(ws, f'B{r}', tenant.get('environment'))
        put(ws, f'C{r}', month.get('month'), 'text')
        part_month = month.get('month') == to_month
        inp(ws, f'D{r}', 1 if tenant.get('environment') == 'prod' and not part_month else 0, 'int')
        for (group, field), letter in col_of.items():
            value = (month.get(group) or {}).get(field)
            if value is not None:
                put(ws, f'{letter}{r}', value)
    last = DATA_ROW + max(len(rows), 1) - 1

    # Per-seat and per-tenant values, one column per profile metric. Blank
    # unless the row is included and the numbers it needs are there.
    agents = col_of[('voice', 'active_agents')]
    helper_col = {}
    for j, (metric, label, unit, nums, per_seat, _) in enumerate(PROFILE):
        letter = get_column_letter(helper_start + j)
        helper_col[metric] = letter
        put(ws, f'{letter}5', 'derived', color=MUTED, size=8)
        header(ws, 6, [f'{label}, {unit}'], start=helper_start + j)
        if nums is None:
            continue
        for r in range(DATA_ROW, last + 1):
            cells = [f'{col_of[n]}{r}' for n in nums]
            checks = [f'$D{r}=1'] + [f'ISNUMBER({c})' for c in cells]
            total = '+'.join(cells)
            if per_seat:
                checks.append(f'ISNUMBER({agents}{r})')
                formula = f'=IF(AND({",".join(checks)}),IF({agents}{r}>0,({total})/{agents}{r},""),"")'
            else:
                formula = f'=IF(AND({",".join(checks)}),{total},"")'
            put(ws, f'{letter}{r}', formula, 'num')

    # The profile: what the model uses, and whether it is measured.
    top = last + 3
    section(ws, top, 'Usage profile: measured where there is data, assumed where there is not', 12)
    header(ws, top + 1, ['Metric', 'Unit', 'n (months)', 'Mean', 'Median', 'P25', 'P90', 'Assumption',
                         'Low used', 'Base used', 'High used', 'Basis'])
    band = a('band')
    for j, (metric, label, unit, nums, _, assumption) in enumerate(PROFILE):
        r = top + 2 + j
        put(ws, f'A{r}', label)
        put(ws, f'B{r}', unit, color=MUTED)
        if nums is None:
            put(ws, f'C{r}', 0, 'int')
            for letter in 'DEFG':
                put(ws, f'{letter}{r}', '')
        else:
            rng = f'${helper_col[metric]}${DATA_ROW}:${helper_col[metric]}${last}'
            put(ws, f'C{r}', f'=COUNT({rng})', 'int')
            put(ws, f'D{r}', f'=IF(C{r}=0,"",AVERAGE({rng}))', 'num')
            put(ws, f'E{r}', f'=IF(C{r}=0,"",MEDIAN({rng}))', 'num')
            put(ws, f'F{r}', f'=IF(C{r}=0,"",PERCENTILE({rng},0.25))', 'num')
            put(ws, f'G{r}', f'=IF(C{r}=0,"",PERCENTILE({rng},0.9))', 'num')
        put(ws, f'H{r}', f'={a(assumption)}', 'num')
        put(ws, f'J{r}', f'=IF(C{r}=0,H{r},D{r})', 'num', bold=True)
        put(ws, f'I{r}', f'=IF(C{r}>=3,F{r},J{r}*(1-{band}))', 'num')
        put(ws, f'K{r}', f'=IF(C{r}>=3,G{r},J{r}*(1+{band}))', 'num')
        put(ws, f'L{r}', f'=IF(C{r}=0,"ASSUMPTION",IF(C{r}<3,"Actual mean, banded","Actual P25 / mean / P90"))',
            bold=True, color='B00020')
        USED[metric] = {'low': f'UsageActuals!$I${r}', 'base': f'UsageActuals!$J${r}',
                        'high': f'UsageActuals!$K${r}'}

    notes_at = top + 2 + len(PROFILE) + 2
    section(ws, notes_at, 'Notes carried in the file', 12)
    for k, note in enumerate(doc.get('notes', [])):
        put(ws, f'A{notes_at + 1 + k}', note, size=9, color=MUTED)

    widths(ws, {'A': 34, 'B': 18, 'C': 12, 'D': 10})
    for c in range(5, last_col + 1):
        ws.column_dimensions[get_column_letter(c)].width = 15
    ws.freeze_panes = f'E{DATA_ROW}'
    return ws


# ── UnitEconomics ───────────────────────────────────────────────────────────

UE = {}         # row name → row on UnitEconomics


def ue(name, case='base'):
    col = {'low': 'C', 'base': 'D', 'high': 'E'}[case]
    return f'UnitEconomics!${col}${UE[name]}'


def host_parts(n):
    """Monthly hosting at `n` tenants: (Forge and AWS, Azure, the one the switch picks)."""
    server, azure_var = 'Hosting!$C$5', 'Hosting!$C$6'
    forge = (f"({a('forge_plan')}*{a('host_alloc')}+{server}*{a('host_alloc')}"
             f"+MAX(0,ROUNDUP({n}/{a('cap_srv')},0)-1)*{server})")
    azure = f"(MAX({a('az_min')},ROUNDUP({n}/{a('az_cap')},0))*{a('az_plan')}+{n}*{azure_var})"
    picked = (f"IF({a('host_mode')}=1,{forge},IF({a('host_mode')}=2,{azure},"
              f"IF({n}>={a('host_switch_n')},{azure},{forge})))")
    return forge, azure, picked


def netlify_at(n):
    return f"{a('nf_site')}*({a('nf_plan')}+MAX(0,{n}*{a('nf_credits')}-{a('nf_incl')})*{a('nf_credit')})"


def github_cost():
    return f"{a('gh_seat')}*{a('gh_seats')}"


def build_unit_economics(wb):
    ws = wb.create_sheet('UnitEconomics')
    title(ws, 'Unit economics',
          'Cost of each tool in use at low, base and high usage. Low and high are P25 and P90 of the actuals, '
          'or plus or minus the band on base when there are under three months.', 6)
    header(ws, 4, ['Line', 'Unit', 'Low', 'Base', 'High', 'Note'])
    state = {'row': 5}

    def sect(text):
        section(ws, state['row'], text, 6)
        state['row'] += 1

    def line(name, label, unit, template, fmt='usd', note='', strong=False):
        r = state['row']
        put(ws, f'A{r}', label, bold=strong)
        put(ws, f'B{r}', unit, color=MUTED)
        for col, case in CASES:
            f = re.sub(r'\{u:(\w+)\}', lambda m: USED[m.group(1)][case], template)
            f = re.sub(r'\{r:(\w+)\}', lambda m: f'{col}{UE[m.group(1)]}', f)
            f = re.sub(r'\{b:(\w+)\}', lambda m: f'$D${UE[m.group(1)]}', f)
            f = re.sub(r'\{(\w+)\}', lambda m: a(m.group(1)), f)
            (key if strong else put)(ws, f'{col}{r}', f'={f}', fmt)
        put(ws, f'F{r}', note, size=9, color=MUTED, wrap=True)
        UE[name] = r
        state['row'] += 1

    sect('Usage, per dialer seat per month')
    line('out_min', 'Billed minutes, PSTN outbound', 'minutes', '{u:out_min}', 'int')
    line('in_min', 'Billed minutes, PSTN inbound', 'minutes', '{u:in_min}', 'int')
    line('browser_min', 'Billed minutes, browser leg', 'minutes', '{u:browser_min}', 'int')
    line('rec_min', 'Recorded minutes', 'minutes', '{u:rec_min}', 'int')
    line('calls', 'Calls', 'calls', '{u:calls}', 'int')
    line('amd_calls', 'Answering-machine-detection calls', 'calls', '{u:amd_calls}', 'int')
    line('cust_min', 'Phone minutes a customer would count', 'minutes', '{r:out_min}+{r:in_min}', 'int',
         'Outbound plus inbound. This is what an included-minutes allowance is measured in.')
    state['row'] += 1

    sect('Usage, per tenant per month')
    line('seg_out', 'Outbound SMS segments', 'segments', '{u:seg_out}*{seats}', 'int', 'Per seat, times the seats.')
    line('seg_in', 'Inbound SMS segments', 'segments', '{u:seg_in}*{seats}', 'int')
    line('ai_convos', 'AI conversations', 'conversations', '{u:ai_convos}', 'int')
    state['row'] += 1

    sect('Twilio Programmable Voice: cost per dialer seat per month')
    line('v_out', 'PSTN outbound', '$', '{r:out_min}*{tw_out}')
    line('v_in', 'PSTN inbound', '$', '{r:in_min}*{tw_in}')
    line('v_browser', 'Browser leg', '$', '{r:browser_min}*{tw_client}')
    line('v_rec', 'Recording', '$', '{r:rec_min}*{tw_rec}')
    line('v_store', 'Recording storage, steady state', '$', '{r:rec_min}*{tw_store}*{retention}', 'usd',
         'Each month of recordings is stored for the retention period. With no retention rule this keeps growing.')
    line('v_amd', 'Answering machine detection', '$', '{r:amd_calls}*{tw_amd}')
    line('v_seat', 'Voice usage per seat', '$',
         '{r:v_out}+{r:v_in}+{r:v_browser}+{r:v_rec}+{r:v_store}+{r:v_amd}', strong=True)
    line('v_cpm', 'Cost per phone minute', '$ / minute', 'IF({r:cust_min}>0,{r:v_seat}/{r:cust_min},0)', 'usd4')
    line('v_nums', 'Phone numbers per tenant', '$', '({seats}*{nums_per_seat}+{nums_shared})*{tw_num}')
    line('v_tenant', 'Voice per tenant', '$', '{seats}*{r:v_seat}+{r:v_nums}', strong=True)
    line('v_seat_all', 'Voice per seat, numbers included', '$', '{r:v_tenant}/{seats}')
    line('over_min', 'Overage price per phone minute', '$ / minute', 'ROUNDUP({over_mult}*{b:v_cpm},3)', 'usd4',
         'One price for every tenant, set from the base cost.')
    state['row'] += 1

    sect('Twilio Messaging and 10DLC: cost per tenant per month')
    line('s_out', 'Outbound segments, carrier fee included', '$', '{r:seg_out}*({tw_sms_out}+{tw_car_out})')
    line('s_in', 'Inbound segments, carrier fee included', '$', '{r:seg_in}*({tw_sms_in}+{tw_car_in})')
    line('s_fixed', 'Campaign fee and texting number', '$', '{tw_camp_mo}+{sms_nums}*{tw_num}')
    line('s_tenant', 'Texting per tenant', '$', '{r:s_out}+{r:s_in}+{r:s_fixed}', strong=True)
    line('s_cps', 'Cost per outbound segment', '$ / segment', '{tw_sms_out}+{tw_car_out}', 'usd4')
    line('over_seg', 'Overage price per segment', '$ / segment', 'ROUNDUP({over_mult}*{r:s_cps},3)', 'usd4')
    line('s_once', 'One-time: brand and campaign registration', '$', '{tw_brand}+{tw_vet}')
    state['row'] += 1

    sect('Zoho CRM: the tenant pays Zoho')
    line('z_seat', 'List price per seat', '$', '{zc_list}')
    line('z_tenant', 'List price per tenant', '$', '{zc_list}*{seats}')
    line('z_margin', 'Kore margin per tenant', '$', '{zoho_margin}*{zc_list}*{seats}')
    line('z_credits', 'API credits a day used by call logging', 'credits',
         '{seats}*{r:calls}/{workdays}*{credits_per_call}', 'int')
    line('z_allow', 'API credits a day allowed', 'credits', '{api_allow}', 'int')
    line('z_util', 'Share of the allowance used', '%', '{r:z_credits}/{r:z_allow}', 'pct',
         'A ceiling, not a dollar cost.')
    state['row'] += 1

    sect('Zoho Voice: the tenant pays Zoho')
    line('zv_seat', 'List price per seat', '$', '{zv_user}')
    line('zv_tenant', 'List price per tenant', '$', '{zv_user}*{seats}')
    line('zv_auto_t', 'With the auto dialer, per tenant (planned)', '$', '({zv_user}+{zv_auto})*{seats}', 'usd',
         'ASSUMPTION: no auto dialer price is published.')
    state['row'] += 1

    sect('Zoho Marketing Automation (planned): cost per tenant if texting ran through it')
    line('zma_tenant', 'Plan, number, campaign and texts', '$', '{zma_plan}+{zma_fixed}+{r:seg_out}*{zma_sms}', 'usd',
         'Counts one text per segment.')
    state['row'] += 1

    sect('AI SMS: cost per tenant per month')
    line('ai_msg', 'Texts the AI sends and receives, on Twilio', '$',
         '{r:ai_convos}*{ai_turns}*({ai_seg_out}*({tw_sms_out}+{tw_car_out})+{ai_seg_in}*({tw_sms_in}+{tw_car_in}))', 'usd',
         'No per-line fee: the AI texts ride the same Twilio messaging as the SMS add-on.')
    line('ai_llm', 'AI model tokens', '$',
         '{r:ai_convos}*{ai_turns}*({tok_in}*{llm_in}+{tok_out}*{llm_out})/1000000')
    line('ai_tenant', 'AI SMS per tenant', '$', '{r:ai_msg}+{r:ai_llm}', strong=True)
    state['row'] += 1

    sect("Platform: hosting, client site, lead lines and GitHub, per tenant at today's tenant count")
    n = a('start_tenants')
    r = state['row']
    for label, name, expr in (
        ('Backend hosting', 'h_host', f'{host_parts(n)[2]}/{n}'),
        ('Client site (Netlify)', 'h_site', f'{netlify_at(n)}/{n}'),
        ('GitHub', 'h_git', f'{github_cost()}/{n}'),
    ):
        r = state['row']
        put(ws, f'A{r}', label)
        put(ws, f'B{r}', '$', color=MUTED)
        for col, _ in CASES:
            put(ws, f'{col}{r}', f'={expr}', 'usd')
        UE[name] = r
        state['row'] += 1
    put(ws, f'F{UE["h_host"]}', 'Falls quickly as tenants are added: see the Hosting sheet.', size=9, color=MUTED)
    put(ws, f'F{UE["h_site"]}', 'Does not fall: each tenant site uses its own credits.', size=9, color=MUTED)
    line('h_lines', 'Lead lines on Twilio', '$', '{pl_twilio}+{pl_numbers}', 'usd',
         "From the first client's bill. Scales with a tenant's lead calls, not with tenant count.")
    line('h_platform', 'Platform per tenant', '$', '{r:h_host}+{r:h_site}+{r:h_git}+{r:h_lines}', strong=True)
    state['row'] += 1

    sect('Cost to serve by product, per tenant that takes it, per month')
    line('c_platform', 'Platform', '$', '{r:h_platform}', strong=True)
    line('c_dialer', 'Dialer', '$', '{r:v_tenant}', strong=True)
    line('c_sms', 'SMS', '$', '{r:s_tenant}', strong=True)
    line('c_comp', 'Compliance', '$', '{comp_cogs}', strong=True)
    line('c_ai', 'AI SMS', '$', '{r:ai_tenant}', strong=True)

    widths(ws, {'A': 50, 'B': 14, 'C': 14, 'D': 14, 'E': 14, 'F': 70})
    ws.freeze_panes = 'C5'
    return ws


# ── VoiceSKUs ───────────────────────────────────────────────────────────────

VS = {}


def build_voice_skus(wb):
    ws = wb.create_sheet('VoiceSKUs')
    title(ws, 'Voice: Zoho Voice (A) against the custom Twilio dialer (B)',
          'A costs Kore nothing and earns the reseller margin. B is priced by Kore and carries the Twilio bill.', 6)
    header(ws, 4, ['Line', 'Unit', 'Low', 'Base', 'High', 'Note'])
    state = {'row': 5}

    def sect(text):
        section(ws, state['row'], text, 6)
        state['row'] += 1

    def line(name, label, unit, template, fmt='usd', note='', strong=False):
        r = state['row']
        put(ws, f'A{r}', label, bold=strong)
        put(ws, f'B{r}', unit, color=MUTED)
        for col, case in CASES:
            f = re.sub(r'\{ue:(\w+)\}', lambda m: ue(m.group(1), case), template)
            f = re.sub(r'\{r:(\w+)\}', lambda m: f'{col}{VS[m.group(1)]}', f)
            f = re.sub(r'\{b:(\w+)\}', lambda m: f'$D${VS[m.group(1)]}', f)
            f = re.sub(r'\{(\w+)\}', lambda m: a(m.group(1)), f)
            (key if strong else put)(ws, f'{col}{r}', f'={f}', fmt)
        put(ws, f'F{r}', note, size=9, color=MUTED, wrap=True)
        VS[name] = r
        state['row'] += 1

    sect('A. Zoho Voice, per seat per month')
    line('a_price', 'What the tenant pays Zoho', '$', '{zv_user}', 'usd',
         'Enterprise Telephony Basic, with unlimited US and Canada calling.')
    line('a_rev', 'What Kore earns', '$', '{zoho_margin}*{zv_user}', 'usd', 'Reseller margin. 0% today.')
    line('a_cost', 'What Kore pays', '$', '0')
    line('a_gp', 'Kore gross profit per seat', '$', '{r:a_rev}-{r:a_cost}', strong=True)
    line('a_auto', 'Tenant price with the auto dialer (planned)', '$', '{zv_user}+{zv_auto}', 'usd',
         'ASSUMPTION: no auto dialer price is published.')
    state['row'] += 1

    sect('B. Custom Twilio dialer, per seat per month')
    line('b_price', 'Price', '$', '{p_dialer}')
    line('b_incl', 'Minutes included', 'minutes', '{incl_min}', 'int')
    line('b_used', 'Phone minutes used', 'minutes', '{ue:cust_min}', 'int')
    line('b_cost', 'Cost, numbers included', '$', '{ue:v_seat_all}')
    line('b_gp_flat', 'Gross profit, no overage charged', '$', '{r:b_price}-{r:b_cost}')
    line('b_gm_flat', 'Gross margin, no overage charged', '%', '{r:b_gp_flat}/{r:b_price}', 'pct', strong=True)
    line('b_over_min', 'Minutes over the allowance', 'minutes', 'MAX(0,{r:b_used}-{r:b_incl})', 'int')
    line('b_over_price', 'Overage price', '$ / minute', '{ue:over_min}', 'usd4')
    line('b_over_rev', 'Overage revenue', '$', '{bill_over}*{r:b_over_min}*{r:b_over_price}')
    line('b_rev', 'Revenue with overage', '$', '{r:b_price}+{r:b_over_rev}')
    line('b_gp', 'Gross profit with overage', '$', '{r:b_rev}-{r:b_cost}')
    line('b_gm', 'Gross margin with overage', '%', '{r:b_gp}/{r:b_rev}', 'pct', strong=True)
    state['row'] += 1

    sect('B against A')
    line('prem', 'What the tenant pays extra for B, per seat', '$', '{r:b_price}-{r:a_price}', 'usd',
         'Before overage. A has no minute cap.')
    line('prem_auto', 'The same, against A with the auto dialer', '$', '{r:b_price}-{r:a_auto}')
    line('kore_gain', 'What Kore earns extra from B, per seat', '$', '{r:b_gp}-{r:a_gp}', strong=True)
    state['row'] += 1

    sect('Breakeven')
    num_seat = f"(({a('seats')}*{a('nums_per_seat')}+{a('nums_shared')})*{a('tw_num')}/{a('seats')})"
    line('be_zero', 'Minutes a seat can use before the flat price loses money', 'minutes',
         f"({{p_dialer}}-{num_seat})/{{ue:v_cpm}}", 'int', 'With no overage charged.')
    line('be_target', 'Minutes a seat can use and still hit the margin target', 'minutes',
         f"({{p_dialer}}*(1-{{target_gm}})-{num_seat})/{{ue:v_cpm}}", 'int')
    line('be_vs_a', 'Minutes at which B earns Kore no more than reselling A', 'minutes',
         f"({{p_dialer}}-{{zoho_margin}}*{{zv_user}}-{num_seat})/{{ue:v_cpm}}", 'int')
    line('be_fixed', 'Fixed monthly cost of running the dialer', '$',
         f"{{ue:h_host}}*{{start_tenants}}", 'usd', 'Backend hosting today.')
    line('be_seats', 'Seats needed to cover that fixed cost', 'seats',
         'IF({r:kore_gain}>0,{r:be_fixed}/{r:kore_gain},"never")', 'num', strong=True)
    state['row'] += 1

    # Margin against minutes used, at the base cost per minute.
    top = state['row']
    section(ws, top, 'Margin by minutes used (base cost per minute)', 6)
    header(ws, top + 1, ['Phone minutes per seat', 'Cost per seat', 'Margin, flat price', 'Overage revenue',
                         'Margin with overage', ''])
    cpm = ue('v_cpm')
    for i, minutes in enumerate((500, 800, 1000, 1200, 1500, 2000, 2500, 3000, 4000)):
        r = top + 2 + i
        inp(ws, f'A{r}', minutes, 'int')
        put(ws, f'B{r}', f'=A{r}*{cpm}+{num_seat}', 'usd')
        put(ws, f'C{r}', f"=({a('p_dialer')}-B{r})/{a('p_dialer')}", 'pct')
        put(ws, f'D{r}', f"={a('bill_over')}*MAX(0,A{r}-{a('incl_min')})*{ue('over_min')}", 'usd')
        put(ws, f'E{r}', f"=({a('p_dialer')}+D{r}-B{r})/({a('p_dialer')}+D{r})", 'pct')

    widths(ws, {'A': 54, 'B': 14, 'C': 16, 'D': 16, 'E': 18, 'F': 62})
    ws.freeze_panes = 'C5'
    return ws


# ── Tiers ───────────────────────────────────────────────────────────────────

T = {}


def tiers(name):
    return f'Tiers!{T[name]}'


def build_tiers(wb):
    ws = wb.create_sheet('Tiers')
    title(ws, 'Tiers: platform fee and add-ons',
          'Revenue and cost per tenant, blended by attach rate. Costs are the base case from UnitEconomics. '
          'Gross margin here is tool cost only: no support, sales or payment fees.', 11)
    header(ws, 4, ['Item', 'Price', 'Unit', 'Units per tenant', 'Attach rate', 'Revenue per tenant that takes it',
                   'Blended revenue per tenant', 'Cost per tenant that takes it', 'Blended cost per tenant',
                   'Gross profit per tenant that takes it', 'Gross margin'])
    seats = a('seats')
    zoho_resale = (f"{a('zoho_margin')}*({a('zc_list')}*{seats}+{a('zv_user')}*{seats}*(1-{a('att_dialer')}))")
    items = [
        ('platform', 'Platform', a('p_platform'), '$ / tenant', '1', '1', ue('c_platform')),
        ('dialer', 'Dialer', a('p_dialer'), '$ / seat', seats, a('att_dialer'), ue('c_dialer')),
        ('dialer_over', 'Dialer overage', ue('over_min'), '$ / minute',
         f"{seats}*MAX(0,{ue('cust_min')}-{a('incl_min')})", f"{a('att_dialer')}*{a('bill_over')}", '0'),
        ('sms', 'SMS', a('p_sms'), '$ / tenant', '1', a('att_sms'), ue('c_sms')),
        ('sms_over', 'SMS overage', ue('over_seg'), '$ / segment',
         f"MAX(0,{ue('seg_out')}-{a('incl_seg')})", f"{a('att_sms')}*{a('bill_over')}", '0'),
        ('comp', 'Compliance', a('p_comp'), '$ / tenant', '1', a('att_comp'), ue('c_comp')),
        ('ai', 'AI SMS', a('p_ai'), '$ / tenant', '1', a('att_ai'), ue('c_ai')),
        ('zoho', 'Zoho resale margin', zoho_resale, '$ / tenant', '1', '1', '0'),
    ]
    first = 5
    for i, (name, label, price, unit, units, attach, cost) in enumerate(items):
        r = first + i
        T[f'row_{name}'] = r
        put(ws, f'A{r}', label)
        put(ws, f'B{r}', f'={price}', 'usd4' if 'over' in name else 'usd')
        put(ws, f'C{r}', unit, color=MUTED)
        put(ws, f'D{r}', f'={units}', 'num')
        put(ws, f'E{r}', f'={attach}', 'pct')
        put(ws, f'F{r}', f'=B{r}*D{r}', 'usd')
        put(ws, f'G{r}', f'=F{r}*E{r}', 'usd')
        put(ws, f'H{r}', f'={cost}', 'usd')
        put(ws, f'I{r}', f'=H{r}*E{r}', 'usd')
        put(ws, f'J{r}', f'=F{r}-H{r}', 'usd')
        put(ws, f'K{r}', f'=IF(F{r}>0,J{r}/F{r},0)', 'pct')
    last = first + len(items) - 1
    row = lambda n: T[f'row_{n}']

    r = last + 2
    section(ws, r, 'Per tenant per month, blended', 11)
    summary = [
        ('arpa', 'Revenue per tenant (ARPA)', f'=SUM(G{first}:G{last})', 'usd'),
        ('arpa_year', 'Revenue per tenant per year', '={arpa}*12', 'usd0'),
        ('arpa_flat', 'Revenue per tenant with no overage', f"={{arpa}}-G{row('dialer_over')}-G{row('sms_over')}", 'usd'),
        ('cost', 'Cost per tenant', f'=SUM(I{first}:I{last})', 'usd'),
        ('gp', 'Gross profit per tenant', '={arpa}-{cost}', 'usd'),
        ('gm', 'Gross margin', '=IF({arpa}>0,{gp}/{arpa},0)', 'pct'),
        ('gm_flat', 'Gross margin with no overage', '=IF({arpa_flat}>0,({arpa_flat}-{cost})/{arpa_flat},0)', 'pct'),
    ]
    for i, (name, label, formula, fmt) in enumerate(summary):
        rr = r + 1 + i
        put(ws, f'A{rr}', label, bold=True)
        formula = re.sub(r'\{(\w+)\}', lambda m: T[m.group(1)], formula)
        key(ws, f'B{rr}', formula, fmt)
        T[name] = f'$B${rr}'

    r = r + len(summary) + 2
    section(ws, r, 'Used by the forecast', 11)
    feeds = [
        ('addon', 'Add-on revenue per tenant, blended',
         f"=G{row('dialer')}+G{row('sms')}+G{row('comp')}+G{row('ai')}", 'usd'),
        ('over', 'Overage and resale revenue per tenant, blended',
         f"=G{row('dialer_over')}+G{row('sms_over')}+G{row('zoho')}", 'usd'),
        ('varcost', 'Cost per tenant, blended, without hosting and client site',
         f"={{cost}}-I{row('platform')}", 'usd'),
        ('once_rev', 'One-time fees per new tenant', f"={a('p_onboard')}+{a('att_sms')}*{a('sms_setup')}", 'usd'),
        ('once_cost', 'One-time cost per new tenant', f"={a('att_sms')}*{ue('s_once')}", 'usd'),
    ]
    for i, (name, label, formula, fmt) in enumerate(feeds):
        rr = r + 1 + i
        put(ws, f'A{rr}', label)
        formula = re.sub(r'\{(\w+)\}', lambda m: T[m.group(1)], formula)
        put(ws, f'B{rr}', formula, fmt, bold=True)
        T[name] = f'$B${rr}'

    r = r + len(feeds) + 2
    section(ws, r, 'Margin by add-on, overage folded in (per tenant that takes it)', 11)
    header(ws, r + 1, ['Add-on', 'Revenue', 'Cost', 'Gross profit', 'Gross margin', 'Gross margin with no overage'])
    addons = [
        ('Platform', f"F{row('platform')}", f"H{row('platform')}", f"F{row('platform')}"),
        ('Dialer', f"F{row('dialer')}+{a('bill_over')}*F{row('dialer_over')}", f"H{row('dialer')}", f"F{row('dialer')}"),
        ('SMS', f"F{row('sms')}+{a('bill_over')}*F{row('sms_over')}", f"H{row('sms')}", f"F{row('sms')}"),
        ('Compliance', f"F{row('comp')}", f"H{row('comp')}", f"F{row('comp')}"),
        ('AI SMS', f"F{row('ai')}", f"H{row('ai')}", f"F{row('ai')}"),
    ]
    for i, (label, rev, cost, flat) in enumerate(addons):
        rr = r + 2 + i
        put(ws, f'A{rr}', label)
        put(ws, f'B{rr}', f'={rev}', 'usd')
        put(ws, f'C{rr}', f'={cost}', 'usd')
        put(ws, f'D{rr}', f'=B{rr}-C{rr}', 'usd')
        key(ws, f'E{rr}', f'=IF(B{rr}>0,D{rr}/B{rr},0)', 'pct')
        put(ws, f'F{rr}', f'=IF({flat}>0,({flat}-C{rr})/({flat}),0)', 'pct')
        T[f'gm_{label}'] = f'$E${rr}'

    r = r + 2 + len(addons) + 1
    section(ws, r, 'Bundles (1 = in the bundle)', 11)
    header(ws, r + 1, ['Item', 'Core', 'Growth', 'Complete'])
    flags = {'platform': (1, 1, 1), 'dialer': (0, 1, 1), 'sms': (0, 1, 1), 'comp': (0, 0, 1), 'ai': (0, 0, 1)}
    labels = {'platform': 'Platform', 'dialer': 'Dialer, all seats', 'sms': 'SMS', 'comp': 'Compliance', 'ai': 'AI SMS'}
    flag_row = {}
    for i, name in enumerate(flags):
        rr = r + 2 + i
        flag_row[name] = rr
        put(ws, f'A{rr}', labels[name])
        for col, flag in zip('BCD', flags[name]):
            inp(ws, f'{col}{rr}', flag, 'int')
    base = r + 2 + len(flags)
    lines = [
        ('Price per month', lambda c: '(' + '+'.join(f"{c}{flag_row[n]}*$F${row(n)}" for n in flags) + f")*(1-{a('bundle_disc')})", 'usd'),
        ('Overage per month at base usage', lambda c: f"{a('bill_over')}*({c}{flag_row['dialer']}*$F${row('dialer_over')}+{c}{flag_row['sms']}*$F${row('sms_over')})", 'usd'),
        ('Cost per month', lambda c: '+'.join(f"{c}{flag_row[n]}*$H${row(n)}" for n in flags), 'usd'),
        ('Gross margin, price only', lambda c: f"IF({c}{base}>0,({c}{base}-{c}{base + 2})/{c}{base},0)", 'pct'),
        ('Gross margin with overage', lambda c: f"IF({c}{base}+{c}{base + 1}>0,({c}{base}+{c}{base + 1}-{c}{base + 2})/({c}{base}+{c}{base + 1}),0)", 'pct'),
    ]
    for i, (label, build, fmt) in enumerate(lines):
        rr = base + i
        put(ws, f'A{rr}', label, bold=True)
        for col in 'BCD':
            key(ws, f'{col}{rr}', f'={build(col)}', fmt)

    widths(ws, {'A': 52, 'B': 16, 'C': 16, 'D': 16, 'E': 16, 'F': 20, 'G': 18, 'H': 20, 'I': 18, 'J': 20, 'K': 14})
    ws.freeze_panes = 'B5'
    return ws


# ── Hosting ─────────────────────────────────────────────────────────────────

def build_hosting(wb):
    ws = wb.create_sheet('Hosting')
    title(ws, 'Hosting: Forge and AWS today, Azure at scale',
          'The switch is host_mode on the Assumptions sheet: 1 Forge and AWS, 2 Azure, 3 Azure from a tenant count.', 12)
    put(ws, 'A5', 'One server, all-in')
    put(ws, 'B5', '$ / month', color=MUTED)
    put(ws, 'C5', f"={a('ec2_inst')}+{a('ebs_gb')}*{a('ebs_price')}+{a('ip_addr')}", 'usd', bold=True)
    put(ws, 'D5', 'Instance, disk and public address.', size=9, color=MUTED)
    put(ws, 'A6', 'Azure cost that grows with each tenant')
    put(ws, 'B6', '$ / tenant / month', color=MUTED)
    put(ws, 'C6', (f"={a('kv_ops')}/10000*{a('kv_price')}+{a('blob_gb')}*{a('blob_price')}"
                   f"+{a('tbl_gb')}*{a('tbl_price')}+{a('tbl_txn')}/10000*{a('tbl_txn_price')}"), 'usd4', bold=True)
    put(ws, 'D6', 'Key Vault operations, Blob and Table storage, Table transactions.', size=9, color=MUTED)
    put(ws, 'A7', 'Switch (host_mode)')
    put(ws, 'C7', f"={a('host_mode')}", 'int', bold=True)
    put(ws, 'D7', f'=CHOOSE(C7,"Forge and AWS","Azure","Azure from "&{a("host_switch_n")}&" tenants")', size=9, color=MUTED)

    header(ws, 9, ['Tenants', 'Servers', 'Forge and AWS, $ / month', 'Per tenant', 'Instances',
                   'Azure, $ / month', 'Per tenant', 'Selected, $ / month', 'Selected per tenant',
                   'Client sites, $ / month', 'Per tenant', 'Hosting and site per tenant'])
    for i, count in enumerate((1, 2, 5, 10, 25, 50, 100, 250, 500)):
        r = 10 + i
        n = f'$A{r}'
        forge, azure, picked = host_parts(n)
        inp(ws, f'A{r}', count, 'int')
        put(ws, f'B{r}', f"=ROUNDUP({n}/{a('cap_srv')},0)", 'int')
        put(ws, f'C{r}', f'={forge}', 'usd')
        put(ws, f'D{r}', f'=C{r}/{n}', 'usd')
        put(ws, f'E{r}', f"=MAX({a('az_min')},ROUNDUP({n}/{a('az_cap')},0))", 'int')
        put(ws, f'F{r}', f'={azure}', 'usd')
        put(ws, f'G{r}', f'=F{r}/{n}', 'usd')
        put(ws, f'H{r}', f'={picked}', 'usd')
        key(ws, f'I{r}', f'=H{r}/{n}', 'usd')
        put(ws, f'J{r}', f'={netlify_at(n)}', 'usd')
        put(ws, f'K{r}', f'=J{r}/{n}', 'usd')
        key(ws, f'L{r}', f'=I{r}+K{r}', 'usd')
    widths(ws, {'A': 40, 'B': 18, 'C': 18, 'D': 14, 'E': 12, 'F': 16, 'G': 14, 'H': 16, 'I': 16, 'J': 18, 'K': 14, 'L': 18})
    return ws


# ── Forecast ────────────────────────────────────────────────────────────────

FC = {}


def build_forecast(wb):
    ws = wb.create_sheet('Forecast')
    title(ws, 'Forecast: 24 months, three scenarios',
          'Tenants ramp to a share of the addressable firms at month 24. Revenue and cost per tenant come from Tiers.', 15)
    scenarios = [('Low', 'pen_low'), ('Base', 'pen_base'), ('High', 'pen_high')]
    cols = ['Month', 'Active tenants', 'New tenants', 'Platform MRR', 'Add-on MRR', 'Overage and resale MRR',
            'Total MRR', 'ARR run rate', 'One-time fees', 'Cost: usage', 'Cost: hosting, sites and lead lines',
            'Cost: one-time', 'Gross profit', 'Gross margin', 'Cumulative revenue']

    header(ws, 4, ['Scenario', 'Share of firms at month 24', 'Addressable firms', 'Tenants, month 12',
                   'MRR, month 12', 'Tenants, month 24', 'MRR, month 24', 'ARR, month 24',
                   'Gross margin, month 24', 'One-time fees, 24 months', 'Revenue, 24 months'])
    block_top = lambda i: 11 + i * 28

    for i, (label, pen) in enumerate(scenarios):
        top = block_top(i)
        section(ws, top, f'{label} scenario', 15)
        header(ws, top + 1, cols)
        start, churn = a('start_tenants'), a('churn_m')
        for m in range(1, 25):
            r = top + 1 + m
            prev = start if m == 1 else f'B{r - 1}'
            put(ws, f'A{r}', m, 'int')
            put(ws, f'B{r}', f"=MAX({start},ROUND({a('addressable')}*{a(pen)}*(A{r}/24)^{a('ramp_exp')},0))", 'int')
            put(ws, f'C{r}', f'=MAX(0,B{r}-{prev}+ROUND({churn}*{prev},0))', 'int')
            put(ws, f'D{r}', f"=B{r}*{a('p_platform')}", 'usd0')
            put(ws, f'E{r}', f"=B{r}*{tiers('addon')}", 'usd0')
            put(ws, f'F{r}', f"=B{r}*{tiers('over')}", 'usd0')
            put(ws, f'G{r}', f'=D{r}+E{r}+F{r}', 'usd0', bold=True)
            put(ws, f'H{r}', f'=G{r}*12', 'usd0')
            put(ws, f'I{r}', f"=C{r}*{tiers('once_rev')}", 'usd0')
            put(ws, f'J{r}', f"=B{r}*{tiers('varcost')}", 'usd0')
            put(ws, f'K{r}', (f'={host_parts(f"B{r}")[2]}+{netlify_at(f"B{r}")}+{github_cost()}'
                              f"+B{r}*({a('pl_twilio')}+{a('pl_numbers')})"), 'usd0')
            put(ws, f'L{r}', f"=C{r}*{tiers('once_cost')}", 'usd0')
            put(ws, f'M{r}', f'=G{r}+I{r}-J{r}-K{r}-L{r}', 'usd0')
            put(ws, f'N{r}', f'=IF(G{r}+I{r}>0,M{r}/(G{r}+I{r}),0)', 'pct')
            put(ws, f'O{r}', f'=G{r}+I{r}' if m == 1 else f'=O{r - 1}+G{r}+I{r}', 'usd0')
        m12, m24, first = top + 13, top + 25, top + 2
        s = 5 + i
        put(ws, f'A{s}', label, bold=True)
        put(ws, f'B{s}', f'={a(pen)}', 'pct')
        put(ws, f'C{s}', f"={a('addressable')}", 'int')
        put(ws, f'D{s}', f'=B{m12}', 'int')
        put(ws, f'E{s}', f'=G{m12}', 'usd0')
        put(ws, f'F{s}', f'=B{m24}', 'int')
        key(ws, f'G{s}', f'=G{m24}', 'usd0')
        key(ws, f'H{s}', f'=H{m24}', 'usd0')
        key(ws, f'I{s}', f'=N{m24}', 'pct')
        put(ws, f'J{s}', f'=SUM(I{first}:I{m24})', 'usd0')
        put(ws, f'K{s}', f'=O{m24}', 'usd0')
        FC[label.lower()] = {'tenants24': f'Forecast!$F${s}', 'mrr24': f'Forecast!$G${s}', 'arr24': f'Forecast!$H${s}'}

    put(ws, 'A9', 'Market case is mkt_case on the Assumptions sheet. Gross margin is tool cost only.',
        italic=True, color=MUTED, size=9)
    widths(ws, {'A': 14, 'B': 16, 'C': 14, 'D': 16, 'E': 16, 'F': 18, 'G': 16, 'H': 16, 'I': 16, 'J': 16,
                'K': 18, 'L': 14, 'M': 16, 'N': 14, 'O': 18})
    ws.freeze_panes = 'B5'
    return ws


# ── Sensitivity ─────────────────────────────────────────────────────────────

def arpa_expr(**over):
    """Blended revenue per tenant per month, with any driver swapped for a test cell."""
    p = lambda name, default: over.get(name, default)
    seats = p('seats', a('seats'))
    att_d, att_s = p('att_d', a('att_dialer')), p('att_s', a('att_sms'))
    att_c, att_ai = p('att_c', a('att_comp')), p('att_ai', a('att_ai'))
    minutes = p('minutes', ue('cust_min'))
    seg_out = p('seg_out', ue('seg_out'))
    bill = a('bill_over')
    return (
        f"({p('platform', a('p_platform'))}"
        f"+{att_d}*{seats}*({a('p_dialer')}+{bill}*MAX(0,{minutes}-{a('incl_min')})*{ue('over_min')})"
        f"+{att_s}*({a('p_sms')}+{bill}*MAX(0,{seg_out}-{a('incl_seg')})*{ue('over_seg')})"
        f"+{att_c}*{a('p_comp')}+{att_ai}*{a('p_ai')}"
        f"+{a('zoho_margin')}*({a('zc_list')}*{seats}+{a('zv_user')}*{seats}*(1-{att_d})))"
    )


def cost_expr(**over):
    """Blended cost per tenant per month, same swaps. Voice cost scales with minutes used."""
    p = lambda name, default: over.get(name, default)
    seats = p('seats', a('seats'))
    att_d, att_s = p('att_d', a('att_dialer')), p('att_s', a('att_sms'))
    att_c, att_ai = p('att_c', a('att_comp')), p('att_ai', a('att_ai'))
    minutes = p('minutes', ue('cust_min'))
    seg_out = p('seg_out', ue('seg_out'))
    numbers = f"({seats}*{a('nums_per_seat')}+{a('nums_shared')})*{a('tw_num')}"
    return (
        f"({ue('c_platform')}"
        f"+{att_d}*({seats}*{minutes}*{ue('v_cpm')}+{numbers})"
        f"+{att_s}*({seg_out}*{ue('s_cps')}+{ue('s_in')}+{ue('s_fixed')})"
        f"+{att_c}*{ue('c_comp')}+{att_ai}*{ue('c_ai')})"
    )


def build_sensitivity(wb):
    ws = wb.create_sheet('Sensitivity')
    title(ws, 'Sensitivity',
          'One driver changes at a time; everything else stays at base. Blue test values can be overwritten.', 5)
    tenants24 = FC['base']['tenants24']
    grids = [
        ('Phone minutes per dialer seat per month', 'minutes', 'int', (600, 1000, 1200, 1500, 2000, 2500, 3000, 4000)),
        ('Seats per tenant', 'seats', 'int', (2, 3, 4, 6, 8, 12)),
        ('Outbound SMS segments per tenant per month', 'seg_out', 'int', (500, 1000, 2500, 5000, 7500, 10000, 20000)),
        ('Dialer attach rate', 'att_d', 'pct', (0.2, 0.4, 0.6, 0.8, 1.0)),
        ('SMS attach rate', 'att_s', 'pct', (0.2, 0.35, 0.5, 0.65, 0.8)),
        ('Compliance attach rate', 'att_c', 'pct', (0.1, 0.25, 0.4, 0.6, 0.8)),
        ('AI SMS attach rate', 'att_ai', 'pct', (0.05, 0.1, 0.2, 0.35, 0.5)),
        ('Platform fee', 'platform', 'usd0', (299, 349, 399, 449, 499)),
        ('Addressable firms', 'firms', 'int', (901, 1256, 4664, 9260)),
    ]
    r = 4
    put(ws, f'A{r}', 'Check: the formula used below, at base, against Tiers', italic=True, color=MUTED, size=9)
    put(ws, f'B{r}', f'={arpa_expr()}', 'usd')
    put(ws, f'C{r}', f"={tiers('arpa')}", 'usd')
    put(ws, f'D{r}', f'={cost_expr()}', 'usd')
    put(ws, f'E{r}', f"={tiers('cost')}", 'usd')
    r += 2
    for label, driver, fmt, values in grids:
        section(ws, r, label, 5)
        header(ws, r + 1, [label, 'Revenue per tenant per month', 'Cost per tenant per month', 'Gross margin',
                           'ARR at month 24, base scenario'])
        for i, value in enumerate(values):
            rr = r + 2 + i
            inp(ws, f'A{rr}', value, fmt)
            over = {} if driver == 'firms' else {driver: f'$A{rr}'}
            put(ws, f'B{rr}', f'={arpa_expr(**over)}', 'usd')
            put(ws, f'C{rr}', f'={cost_expr(**over)}', 'usd')
            put(ws, f'D{rr}', f'=IF(B{rr}>0,(B{rr}-C{rr})/B{rr},0)', 'pct')
            tenants = (f"MAX({a('start_tenants')},ROUND($A{rr}*{a('pen_base')},0))" if driver == 'firms' else tenants24)
            key(ws, f'E{rr}', f'={tenants}*B{rr}*12', 'usd0')
        r += 2 + len(values) + 1
    put(ws, f'A{r}', 'Seats: texting volume is held at base, though more seats would send more. '
                     'Minutes: voice cost moves in step with minutes at the base cost per minute. '
                     'Hosting per tenant is held at today\'s tenant count.', italic=True, color=MUTED, size=9)
    widths(ws, {'A': 46, 'B': 26, 'C': 24, 'D': 16, 'E': 30})
    return ws


def apply_overrides(pairs):
    """`key=value` pairs from --set, applied to the inputs before the sheets are written."""
    inputs = {e['key']: e for _, entries in SECTIONS for e in entries if e['kind'] == 'input'}
    for pair in pairs:
        name, _, raw = pair.partition('=')
        if name not in inputs:
            raise SystemExit(f'--set {pair}: "{name}" is not an input (see assumptions.py)')
        try:
            value = float(raw)
        except ValueError:
            raise SystemExit(f'--set {pair}: "{raw}" is not a number')
        inputs[name]['value'] = int(value) if value.is_integer() and '.' not in raw else value
        inputs[name]['source'] = f"Set on the command line (was: {inputs[name]['source']})"


def main():
    global ACTUALS, OUT
    parser = argparse.ArgumentParser(description='Build the HaulOS pricing model workbook.')
    parser.add_argument('--actuals', type=Path, default=ACTUALS, help='usage-actuals.json to read')
    parser.add_argument('--out', type=Path, default=OUT, help='workbook to write')
    parser.add_argument('--set', action='append', default=[], metavar='KEY=VALUE', help='override one input')
    args = parser.parse_args()
    ACTUALS, OUT = args.actuals.resolve(), args.out.resolve()
    apply_overrides(args.set)

    assign_rows()
    wb = Workbook()
    build_assumptions(wb)
    build_usage_actuals(wb)
    build_unit_economics(wb)
    build_voice_skus(wb)
    build_tiers(wb)
    build_hosting(wb)
    build_forecast(wb)
    build_sensitivity(wb)
    wb.calculation.fullCalcOnLoad = True
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    shown = OUT.relative_to(ROOT) if OUT.is_relative_to(ROOT) else OUT
    print(f'wrote {shown} ({len(wb.sheetnames)} sheets: {", ".join(wb.sheetnames)})')


if __name__ == '__main__':
    main()
